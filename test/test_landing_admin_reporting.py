"""Isolated factual reporting/export checks. No production network or databases."""

import ast
import csv
from contextlib import closing
from datetime import datetime, timedelta
import importlib.util
from io import BytesIO, StringIO
import operator
import os
from pathlib import Path
import sqlite3
import sys
import time
from types import ModuleType, SimpleNamespace
from urllib.parse import parse_qs, urlsplit
from uuid import uuid4

from flask import Flask
import openpyxl
import peewee as pw
from playhouse.sqlite_ext import JSONField
import pytest

# The host Python has an unrelated preinstalled requests/chardet mismatch.
# Transport calls below are mocked and never use that dependency combination.
pytestmark = pytest.mark.filterwarnings("ignore:urllib3.*:Warning")

ROOT = Path(__file__).resolve().parents[1]
SITE = ROOT / "new site for swipies"
sys.path.insert(0, str(SITE))


def load(path, name):
    spec = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(spec)
    sys.modules[name] = module
    spec.loader.exec_module(module)
    return module


core = load(ROOT / "api/analytics_reporting.py", "api.analytics_reporting")
adapter = load(SITE / "landing_admin_api.py", "landing_admin_api")
providers = load(SITE / "landing_analytics_providers.py", "landing_analytics_providers")
BASE = {
    "start": "2025-06-02T00:00:00Z",
    "end": "2025-06-04T00:00:00Z",
    "timezone": "UTC",
    "group": "day",
}


@pytest.mark.parametrize(
    "group,count", [("hour", 48), ("day", 2), ("week", 1), ("month", 1)]
)
def test_calendar_groups(group, count):
    f = core.Filters.parse({**BASE, "group": group})
    assert len(core.buckets(f.current, f)) == count
    assert f.previous.end == f.current.start
    assert f.current.end - f.current.start == f.previous.end - f.previous.start


def test_dst_repeated_hour_and_calendar_months():
    f = core.Filters.parse(
        {
            **BASE,
            "start": "2025-10-25T22:00:00Z",
            "end": "2025-10-26T23:00:00Z",
            "timezone": "Europe/Berlin",
            "group": "hour",
        }
    )
    buckets = core.buckets(f.current, f)
    assert len(buckets) == 25 and len({b["start"] for b in buckets}) == 25
    with pytest.raises(ValueError, match="повторяется"):
        providers._native_index(core.blank_period(f.current, f), f)
    f = core.Filters.parse(
        {
            **BASE,
            "start": "2024-01-31T00:00:00Z",
            "end": "2024-03-02T00:00:00Z",
            "group": "month",
        }
    )
    assert [b["start"] for b in core.buckets(f.current, f)] == [
        "2024-01-31T00:00:00Z",
        "2024-02-01T00:00:00Z",
        "2024-03-01T00:00:00Z",
    ]


@pytest.mark.parametrize(
    "change",
    [
        {"end": "2025-06-02T00:00:00Z"},
        {"start": "2025-06-02T00:01:00Z"},
        {"group": "year"},
        {"timezone": "Not/AZone"},
        {"comparison": "custom"},
        {"end": "2099-01-01T00:00:00Z"},
        {"start": "2020-01-01T00:00:00Z"},
        {"utm_source": "x" * 129},
    ],
)
def test_invalid_filters(change):
    with pytest.raises(ValueError):
        core.Filters.parse({**BASE, **change})


def test_local_whole_hour_in_half_hour_timezone():
    f = core.Filters.parse(
        {
            **BASE,
            "start": "2025-06-01T18:30:00Z",
            "end": "2025-06-02T18:30:00Z",
            "timezone": "Asia/Kolkata",
        }
    )
    assert len(core.buckets(f.current, f)) == 1


def test_period_distincts_boundaries_and_attribution():
    f = core.Filters.parse(BASE)
    visitors = [
        {
            "id": i,
            "visitor_id": "same",
            "ip": "1.2.3.4",
            "create_date": "2025-06-02T12:00:00Z",
            "session_duration": 10,
            "utm_source": "Телеграм",
            "utm_campaign": "launch",
        }
        for i in range(600)
    ]
    visitors += [
        {**visitors[0], "id": 600, "create_date": "2025-06-03T01:00:00Z"},
        {**visitors[0], "id": 601, "create_date": BASE["end"]},
    ]
    p = core.traffic_period(f.current, f, visitors, [])
    assert p["totals"]["visits"] == 601 and p["totals"]["visitors"] == 1
    assert sum(b["values"]["visitors"] for b in p["buckets"]) == 2
    assert p["totals"]["unique_ips"] == 1 and p["campaigns"][0]["visits"] == 601
    f = core.Filters.parse({**BASE, "utm_source": "телеграм"})
    p = core.traffic_period(
        f.current, f, visitors, [{"id": 1, "create_date": BASE["start"]}]
    )
    assert (
        p["totals"]["visits"] == 601
        and "leads" not in p["totals"]
        and p["unavailable"]["leads"]
    )


def test_deltas_no_infinity_and_percentage_points():
    assert core.delta(12, 0)["state"] == "new" and core.delta(12, 0)["percent"] is None
    assert core.delta(0, 0)["percent"] == 0
    assert core.delta(5, None)["absolute"] is None
    assert core.delta(20, 10, "percent") == {
        "absolute": 10,
        "percent": 100,
        "state": "compared",
        "unit": "percentage_points",
    }


def test_paid_only_revenue_distinct_payers_and_missing_amount():
    f = core.Filters.parse(BASE)
    payments = [
        {
            "create_date": BASE["start"],
            "user_id": "one",
            "status": status,
            "amount_uzs": amount,
        }
        for status, amount in [("PAID", 100), ("paid", 200), ("FAILED", 900)]
    ]
    p = core.platform_period(f.current, f, [], payments)
    assert p["totals"]["revenue_uzs"] == 300 and p["totals"]["paying_users"] == 1
    assert p["totals"]["paid_payments"] == 2 and p["totals"]["avg_payment_uzs"] == 150
    payments[0]["amount_uzs"] = None
    p = core.platform_period(f.current, f, [], payments)
    assert p["totals"]["revenue_uzs"] is None and p["unavailable"]["revenue_uzs"]


def test_workbook_numeric_values_safe_text_and_metadata():
    f = core.Filters.parse(BASE)
    p = core.traffic_period(f.current, f, [], [])
    report = core.assemble(f, p, core.traffic_period(f.previous, f, [], []), "landing")
    book = openpyxl.load_workbook(BytesIO(core.xlsx_bytes(core.report_sheets(report))))
    assert book.sheetnames == [
        "Параметры",
        "Сравнение",
        "Текущий период",
        "Период сравнения",
    ]
    metadata = dict(book["Параметры"].values)
    assert metadata["timezone"] == "UTC" and metadata["equal_duration"] is True
    assert (
        book["Сравнение"]["B2"].value == 0 and book["Сравнение"]["B2"].data_type == "n"
    )
    literal = openpyxl.load_workbook(
        BytesIO(core.xlsx_bytes([("Тест", [["Значение"], ["=1+1"], [True]])]))
    )
    assert (
        literal.active["A2"].value == "=1+1" and literal.active["A2"].data_type == "s"
    )
    assert literal.active["A3"].value is True
    output = core.csv_bytes([["a", "b"], ["=1+1", 'line\n"quoted"']]).decode(
        "utf-8-sig"
    )
    assert list(csv.reader(StringIO(output)))[1] == ["'=1+1", 'line\n"quoted"']


@pytest.fixture
def web(tmp_path):
    path = tmp_path / "journal.sqlite"
    with closing(sqlite3.connect(path)) as db:
        db.execute(
            "CREATE TABLE visitors (id INTEGER PRIMARY KEY, visitor_id TEXT, ip TEXT, country TEXT, city TEXT, isp TEXT, device_type TEXT, browser TEXT, language TEXT, session_duration INTEGER, scroll_depth INTEGER, utm_source TEXT, utm_medium TEXT, utm_campaign TEXT, cookie_consent TEXT, page_url TEXT, referrer TEXT, create_date TEXT)"
        )
        db.execute(
            "CREATE TABLE leads (id INTEGER PRIMARY KEY, name TEXT, company TEXT, email TEXT, phone TEXT, message TEXT, status TEXT, referral_code TEXT, create_date TEXT)"
        )
        db.executemany(
            "INSERT INTO visitors (id,visitor_id,ip,utm_source,utm_campaign,create_date,session_duration) VALUES (?,?,?,?,?,?,?)",
            [
                (
                    i,
                    "v" + str(i % 20),
                    "127.0.0.1",
                    "Телеграм" if i % 2 == 0 else "",
                    "launch",
                    "2025-06-02T01:00:00Z",
                    i,
                )
                for i in range(605)
            ],
        )
        db.execute(
            'INSERT INTO visitors (id,visitor_id,create_date) VALUES (9999,"outside","2025-06-04T00:00:00Z")'
        )
        db.execute(
            'INSERT INTO leads VALUES (1,"=2+2","Shop","test@example.invalid","123","line1\nline2","1","ref","2025-06-02 02:00:00")'
        )
        db.commit()
    calls = []
    flags = {"platform": True, "superuser": True, "otp": False}

    def upstream(path, authorization="", method="GET", body=None):
        calls.append((path, method, body))
        if path == "/api/v1/users/me":
            return 200, {"code": 0, "data": {"is_superuser": flags["superuser"]}}, {}
        if path == "/api/v1/auth/login":
            if flags["otp"]:
                return 200, {"code": 0, "data": {"requires_2fa": True}}, {}
            return 200, {"code": 0}, {"Authorization": "preview-test-session"}
        if path == "/api/v1/auth/activate":
            return 200, {"code": 0}, {"Authorization": "preview-test-session"}
        if "/periods?" in path:
            if not flags["platform"]:
                return 503, {}, {}
            params = {k: v[0] for k, v in parse_qs(urlsplit(path).query).items()}
            f = core.Filters.parse(params)
            return (
                200,
                {
                    "code": 0,
                    "data": core.assemble(
                        f,
                        core.platform_period(
                            f.current,
                            f,
                            [{"id": "one", "create_date": BASE["start"]}],
                            [],
                        ),
                        core.platform_period(f.previous, f, [], []),
                        "platform",
                    ),
                },
                {},
            )
        if "/users?" in path:
            params = parse_qs(urlsplit(path).query)
            page = int(params.get("page", [1])[0])
            allrows = [
                {"id": str(i), "email": f"u{i}@example.invalid"} for i in range(205)
            ]
            return (
                200,
                {
                    "code": 0,
                    "data": {
                        "data": allrows[(page - 1) * 100 : page * 100],
                        "total": 205,
                        "page": page,
                        "page_size": 100,
                    },
                },
                {},
            )
        raise AssertionError(path)

    app = Flask(__name__)
    app.testing = True
    adapter.register_admin_analytics(app, path, upstream)
    return SimpleNamespace(
        client=app.test_client(), flags=flags, calls=calls, path=path
    )


def post(web, path, **values):
    return web.client.post(
        "/api/admin/analytics/" + path,
        headers={"Authorization": "Bearer preview-test-session"},
        json={"filters": BASE, **values},
    )


def test_auth_requires_live_superuser_not_local_storage(web):
    response = web.client.post("/api/admin/analytics/report", json={"filters": BASE})
    assert response.status_code == 401
    web.flags["superuser"] = False
    assert post(web, "report").status_code == 403
    assert post(web, "export").status_code == 403
    assert post(web, "records").status_code == 403


def test_login_uses_standard_auth_and_otp_verification(web):
    web.flags["otp"] = True
    response = post(
        web, "login", email="admin@example.invalid", password="encrypted-test"
    )
    assert (
        response.json["data"]["requires_2fa"]
        and "authorization" not in response.json["data"]
    )
    response = post(web, "verify", email="admin@example.invalid", code="123456")
    assert response.json["data"]["authorization"].startswith("Bearer ")
    assert any(call[0] == "/api/v1/auth/login" for call in web.calls)
    assert any(call[0] == "/api/v1/auth/activate" for call in web.calls)
    assert not any("/admin/login" in call[0] for call in web.calls)


def test_report_complete_and_platform_unavailable_is_not_zero(web):
    response = post(web, "report")
    assert response.status_code == 200, response.json
    data = response.json["data"]
    assert data["current"]["totals"]["visits"] == 605
    assert data["current"]["totals"]["new_users"] == 1
    assert response.headers["Cache-Control"] == "no-store"
    web.flags["platform"] = False
    data = post(web, "report").json["data"]
    metric = next(x for x in data["metrics"] if x["key"] == "new_users")
    assert metric["current"] is None and metric["reason"]


def test_records_pages_unicode_filters_and_literal_search(web):
    response = post(web, "records", table="visitors", page=6)
    assert response.status_code == 200, response.json
    assert response.json["total"] == 605 and len(response.json["data"]) == 100
    response = post(
        web, "records", table="campaigns", filters={**BASE, "utm_source": "телеграм"}
    )
    assert response.json["total"] == 303
    assert post(web, "records", table="visitors", search="%").json["total"] == 0
    assert (
        post(
            web, "records", table="leads", filters={**BASE, "utm_source": "telegram"}
        ).status_code
        == 400
    )


@pytest.mark.parametrize(
    "kind,expected",
    [("visitors", 606), ("campaigns", 606), ("leads", 2), ("users", 206)],
)
def test_export_all_filtered_records_not_visible_page(web, kind, expected):
    response = post(
        web,
        "export",
        kind=kind,
        format="xlsx",
        credentials={"oauth_token": "never-export-this"},
    )
    assert response.status_code == 200, response.json
    workbook = openpyxl.load_workbook(BytesIO(response.data))
    assert workbook["Записи"].max_row == expected
    assert "never-export-this" not in str(list(workbook["Параметры"].values))
    assert dict(workbook["Параметры"].values)["start"] == BASE["start"]
    if kind == "leads":
        assert workbook["Записи"]["B2"].data_type == "s"


def test_report_export_csv_filters_search_and_readonly(web):
    before = web.path.read_bytes()
    response = post(web, "export", format="xlsx")
    assert response.status_code == 200, response.json
    assert len(openpyxl.load_workbook(BytesIO(response.data)).sheetnames) == 4
    response = post(
        web,
        "export",
        kind="visitors",
        format="csv",
        search="v19",
        filters={**BASE, "utm_campaign": "launch"},
    )
    rows = list(csv.reader(StringIO(response.data.decode("utf-8-sig"))))
    assert (
        len(rows) == 31
        and "timezone" in rows[0]
        and rows[1][rows[0].index("utm_campaign", 17)] == "launch"
    )
    assert rows[1][rows[0].index("search")] == "v19"
    assert web.path.read_bytes() == before


@pytest.mark.parametrize("value", [[], None, "string"])
def test_json_object_validation(web, value):
    response = web.client.post(
        "/api/admin/analytics/report",
        headers={"Authorization": "Bearer preview-test-session"},
        json=value,
        content_type="application/json",
    )
    assert response.status_code == 400


@pytest.mark.parametrize("group", ["hour", "day", "week", "month"])
@pytest.mark.parametrize("intervals", [True, False])
def test_yandex_native_timeline_including_first_day(monkeypatch, group, intervals):
    f = core.Filters.parse({**BASE, "group": group})
    requests = []

    def get(url, **kwargs):
        params = kwargs["params"]
        requests.append(params)
        window = core.Window(
            core.instant(params["date1"]),
            core.instant(params["date2"]) + timedelta(days=1),
        )
        native = core.buckets(window, f)
        if "bytime" in url:
            result = {
                "data": [
                    {
                        "metrics": [
                            [1] * len(native),
                            [2] * len(native),
                            [3] * len(native),
                            [10] * len(native),
                            [60] * len(native),
                            [1] * len(native),
                        ]
                    }
                ]
            }
            if intervals:
                result["time_intervals"] = [
                    [
                        b["label"].split("T")[0] if group != "hour" else b["label"],
                        b["end"],
                    ]
                    for b in native
                ]
        elif "dimensions" in params:
            result = {"data": [{"dimensions": [{"name": "Test"}], "metrics": [5]}]}
        else:
            result = {"totals": [1, 4, 6, 10, 60, 2]}
        return SimpleNamespace(status_code=200, json=lambda: result)

    monkeypatch.setattr("requests.get", get)
    a, b, _ = providers.yandex_periods(
        {"counter_id": "123", "oauth_token": "fake-test"}, f, core
    )
    assert a["totals"]["visitors"] == 1
    assert all(bucket["values"]["visits"] == 2 for bucket in a["buckets"])
    assert (
        requests[0]["timezone"] == "+00:00"
        and "ym:s:dateTime>=" in requests[0]["filters"]
    )
    assert "demo_clicks" not in a["totals"] and a["unavailable"]["demo_clicks"]


@pytest.mark.parametrize("group", ["hour", "day", "week", "month"])
def test_ga_totals_not_summed_from_buckets(monkeypatch, group):
    f = core.Filters.parse({**BASE, "group": group})
    import google.oauth2.service_account

    monkeypatch.setattr(
        google.oauth2.service_account.Credentials,
        "from_service_account_info",
        lambda *a, **k: SimpleNamespace(token="fake-test", refresh=lambda _: None),
    )
    requests = []

    def post_(url, **kwargs):
        results = []
        for query in kwargs["json"]["requests"]:
            requests.append(query)
            start = query["dimensionFilter"]
            if "eventName" in str(start):
                start = start["andGroup"]["expressions"][0]
            datevalue = start["andGroup"]["expressions"][0]["filter"]["numericFilter"][
                "value"
            ]["int64Value"]
            window = f.current if datevalue.startswith("20250602") else f.previous
            period = core.blank_period(window, f)
            indexed = providers._native_index(period, f)
            dimensions = [d["name"] for d in query["dimensions"]]
            if dimensions == []:
                rows = [
                    {
                        "metricValues": [
                            {"value": str(v)} for v in [1, 4, 6, 0.25, 60, 2]
                        ]
                    }
                ]
            elif "eventName" in dimensions:
                rows = []
            elif dimensions[0] in ["dateHour", "date", "isoYearIsoWeek", "yearMonth"]:
                rows = [
                    {
                        "dimensionValues": [{"value": key}],
                        "metricValues": [
                            {"value": str(v)} for v in [1, 2, 3, 0.5, 30, 1]
                        ],
                    }
                    for key in indexed
                ]
            else:
                rows = []
            results.append({"rows": rows, "metadata": {"timeZone": "UTC"}})
        return SimpleNamespace(status_code=200, json=lambda: {"reports": results})

    monkeypatch.setattr("requests.post", post_)
    a, b, _ = providers.google_periods(
        {"property_id": "123", "service_account_key": {}}, f, core
    )
    assert a["totals"]["visitors"] == 1 and a["totals"]["bounce_rate"] == 25
    assert all(bucket["values"]["bounce_rate"] == 50 for bucket in a["buckets"])
    assert "numericFilter" in str(requests[0]["dimensionFilter"])


@pytest.fixture
def models(tmp_path, monkeypatch):
    from quart_auth import AuthUser

    db = pw.SqliteDatabase(tmp_path / "platform.sqlite")
    namespace = {
        **vars(pw),
        "DB": db,
        "os": os,
        "operator": operator,
        "AuthUser": AuthUser,
        "EmptyStringCharField": pw.CharField,
        "JSONField": JSONField,
        "get_uuid": lambda: uuid4().hex,
        "settings": SimpleNamespace(DATABASE_TYPE="mysql"),
        "current_timestamp": lambda: int(time.time() * 1000),
        "timestamp_to_date": lambda t: datetime.fromtimestamp(t / 1000),
        "AUTO_DATE_TIMESTAMP_FIELD_PREFIX": {"create", "update"},
    }
    tree = ast.parse((ROOT / "api/db/db_models.py").read_text(encoding="utf-8"))
    names = {"BaseModel", "DataBaseModel", "User", "PaymentTransaction", "PaymentOrder"}
    nodes = [node for node in tree.body if getattr(node, "name", None) in names]
    exec(
        compile(
            ast.fix_missing_locations(
                ast.Module(
                    body=ast.parse("from __future__ import annotations").body + nodes,
                    type_ignores=[],
                )
            ),
            str(ROOT / "api/db/db_models.py"),
            "exec",
        ),
        namespace,
    )
    classes = [
        namespace[name] for name in ["User", "PaymentTransaction", "PaymentOrder"]
    ]
    db.create_tables(classes)
    stub = ModuleType("api.db.db_models")
    stub.DB = db
    for model in classes:
        setattr(stub, model.__name__, model)
    monkeypatch.setitem(sys.modules, "api.db.db_models", stub)
    service = load(
        ROOT / "api/db/services/landing_analytics_service.py",
        "isolated_landing_service",
    )
    yield SimpleNamespace(
        db=db, User=classes[0], Ledger=classes[1], Order=classes[2], service=service
    )
    db.close()


def test_real_models_deduplicate_mirrors_and_exclude_ads(models):
    stamp = int(core.instant(BASE["start"]).timestamp() * 1000)
    models.User.create(
        id="u",
        email="test@example.invalid",
        nickname="Test",
        create_time=stamp,
        utm_source="Телеграм",
        utm_campaign="launch",
    )
    models.User.update(create_time=stamp).where(models.User.id == "u").execute()

    def ledger(id, tx, when, amount):
        models.Ledger.create(
            id=id,
            transaction_id=tx,
            user_id="u",
            tenant_id="u",
            account_email="test@example.invalid",
            plan_type="pro",
            expected_amount_uzs=amount,
            paid_amount_uzs=amount,
            status="PAID",
            create_time=when,
        )
        models.Ledger.update(create_time=when).where(models.Ledger.id == id).execute()

    ledger("l1", "t1", stamp, 100)
    ledger("outside", "t-outside", stamp - 10 * 86400000, 900)

    def order(id, tx, amount, purpose="subscription_upgrade"):
        models.Order.create(
            id=id,
            external_transaction_id=tx,
            user_id="u",
            tenant_id="u",
            purpose=purpose,
            status="paid",
            amount_uzs=amount,
            create_time=stamp,
        )
        models.Order.update(create_time=stamp).where(models.Order.id == id).execute()

    order("mirror", "t1", 100)
    order("late-mirror", "t-outside", 900)
    order("unique", "t2", 200)
    order("duplicate", "t2", 200)
    order("ads", "t3", 700, "advertiser_deposit")
    result = models.service.report(BASE)
    assert result["current"]["totals"]["revenue_uzs"] == 300
    assert result["current"]["totals"]["paid_payments"] == 2
    assert result["current"]["totals"]["new_users"] == 1
    users = models.service.users_report(
        {**BASE, "utm_source": "телеграм", "search": "TEST"}
    )
    assert (
        users["total"] == 1
        and "password" not in users["data"][0]
        and "access_token" not in users["data"][0]
    )


def test_new_routes_require_superuser_and_no_cache():
    tree = ast.parse(
        (ROOT / "api/apps/restful_apis/system_api.py").read_text(encoding="utf-8")
    )
    for name in ["admin_landing_periods_report", "admin_landing_users_report"]:
        node = next(n for n in tree.body if getattr(n, "name", None) == name)
        assert any(
            isinstance(d, ast.Name) and d.id == "login_required"
            for d in node.decorator_list
        )
        source = ast.unparse(node)
        assert "current_user.is_superuser" in source and "'no-store'" in source
